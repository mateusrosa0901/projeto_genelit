"""Conferência com as planilhas calculadas à mão (pasta 'referencias'). Elas NÃO entram no cálculo:
servem só para mostrar, linha a linha, onde o resultado do programa bate ou não com o manual."""
import re

import numpy as np
import pandas as pd
from openpyxl import load_workbook

from .parceiros import PAGOS, RPA
from .util import ap, arred1, nome_base, norm, txt_int, vazio

COMUNS = ["MATRÍCULA", "NOME (ALUNO)", "PARCELA", "NOME PARCEIRO", "DATA PAGTO", "VALOR PAGTO"]


def chave_linha(d):
    """Identifica o pagamento: matrícula + parcela + data + valor pago + curso."""
    def txt(v):
        return "" if vazio(v) else str(txt_int(v)).strip().upper()
    data = pd.to_datetime(d["DATA PAGTO"], errors="coerce").dt.strftime("%Y-%m-%d").fillna("")
    valor = pd.to_numeric(d["VALOR PAGTO"], errors="coerce").round(2).map(lambda v: f"{v:.2f}")
    curso = ap(d["CURSO"], lambda v: norm(v) or "") if "CURSO" in d else ""
    return (ap(d["MATRÍCULA"], txt) + "|" + ap(d["PARCELA"], txt) + "|" + data.astype(str) + "|" +
            valor.astype(str) + "|" + curso)


def _ler_referencia(arq):
    """DataFrame com os valores (fórmulas já calculadas pelo Excel/Google Planilhas). Se o arquivo não tiver
    os valores salvos, calcula as fórmulas simples do VALOR DO PARCEIRO (=K2*0.2, =K2)."""
    df = pd.read_excel(arq)
    df.columns = [re.sub(r"\s+", " ", str(c)).strip().upper() for c in df.columns]
    if "VALOR DO PARCEIRO" in df.columns and df["VALOR DO PARCEIRO"].isna().all():
        wb = load_workbook(arq)
        ws = wb.worksheets[0]
        cab = [re.sub(r"\s+", " ", str(c.value)).strip().upper() for c in ws[1]]
        if "VALOR DO PARCEIRO" in cab and "VALOR PAGTO" in cab:
            cp, ck = cab.index("VALOR DO PARCEIRO") + 1, cab.index("VALOR PAGTO") + 1
            vals = []
            for r in range(2, len(df) + 2):
                f, k = ws.cell(r, cp).value, ws.cell(r, ck).value
                m = re.fullmatch(r"=\s*K\d+\s*(?:\*\s*([0-9.,]+))?", str(f)) if isinstance(f, str) else None
                if m and isinstance(k, (int, float)):
                    vals.append(k * float((m.group(1) or "1").replace(",", ".")))
                else:
                    vals.append(f if isinstance(f, (int, float)) else np.nan)
            df["VALOR DO PARCEIRO"] = vals
    return df


def _tipo_referencia(df):
    cols = set(df.columns)
    if {"NOME PARCEIRO", "VALOR", "PF / PJ"} <= cols:
        return "PLANILHA GERAL"
    if {"MATRÍCULA", "NOME (ALUNO)", "PARCELA", "DATA PAGTO", "VALOR PAGTO"} <= cols:
        polo = pd.to_numeric(df.get("VALOR DO POLO"), errors="coerce")
        parc = pd.to_numeric(df.get("VALOR DO PARCEIRO"), errors="coerce")
        if parc is not None and parc.notna().sum() > 0 and (polo is None or polo.notna().sum() == 0):
            return "PARCEIRO"
        if polo is not None and polo.notna().sum() > 0:
            return "POLO"
    return None


def arquivos_referencia(pasta):
    if not pasta.exists():
        return []
    return sorted(p for p in pasta.rglob("*.xls*") if not p.name.startswith("~$"))


# ======================================== PARCEIROS ========================================
def conferir_parceiros(r, pasta):
    """Compara os relatórios de parceiro e a PLANILHA GERAL feitos à mão com o cálculo.
    Devolve (resumo, linhas) ou (None, None) se não houver referências."""
    dados, cad, cruz, resumo = r["dados"], r["cadastro"], r["cruzador"], r["resumo_parceiros"]
    total_id = dict(zip(resumo["_id"], resumo["TOTAL SEM ARREDONDAR"])) if len(resumo) else {}
    valor_planilha = r["valores_planilha"]
    res, todas = [], []
    for arq in arquivos_referencia(pasta):
        try:
            ref = _ler_referencia(arq)
        except Exception as erro:
            print(f"[Conferência] Não foi possível ler {arq.name}: {erro}")
            continue
        tipo = _tipo_referencia(ref)
        if tipo == "PARCEIRO":
            ref = ref[ref["MATRÍCULA"].notna() & ref["NOME (ALUNO)"].notna()].copy()
            nome = ref["NOME PARCEIRO"].dropna().mode()
            # parceiro = o mesmo ligado às linhas do Pincel com esses nomes (usa nome + CPF/CNPJ)
            ids = dados.loc[dados["NOME PARCEIRO"].isin(set(ref["NOME PARCEIRO"].dropna())), "_parceiro_id"].dropna()
            pid = ids.mode().iloc[0] if len(ids) else (cruz.cruzar(nome.iloc[0], None)[0] if len(nome) else None)
            vazio_df = dados.iloc[0:0].copy()
            calc = dados[(dados["_parceiro_id"] == pid) & dados["PARCEIRO PAGO"]].copy() if pid else vazio_df.copy()
            base = dados[dados["_parceiro_id"] == pid].copy() if pid else vazio_df.copy()
            for d in (ref, calc, base):
                d["_k"] = chave_linha(d)
                d["_n"] = d.groupby("_k").cumcount()
            explica = dict(zip(zip(base["_k"], base["_n"]), base["CLASSIFICAÇÃO PARCEIRO"] + " " +
                               base["MOTIVO PARCEIRO"].fillna("").astype(str)))
            lado_ref = ref[["_k", "_n"] + [c for c in COMUNS + ["CURSO"] if c in ref.columns] +
                           ["VALOR DO PARCEIRO"]].rename(columns={"VALOR DO PARCEIRO": "VALOR MANUAL"})
            lado_calc = calc[["_k", "_n"] + COMUNS + ["CURSO", "% PARCEIRO", "VALOR DO PARCEIRO CALCULADO",
                                                     "EXCEÇÃO PARCEIRO"]].rename(
                columns={"VALOR DO PARCEIRO CALCULADO": "VALOR CALCULADO"})
            j = lado_ref.merge(lado_calc, on=["_k", "_n"], how="outer", suffixes=("", " (calc)"), indicator=True)
            for c in COMUNS + ["CURSO"]:
                if f"{c} (calc)" in j.columns:
                    j[c] = j[c].where(j[c].notna(), j[f"{c} (calc)"])
            manual = pd.to_numeric(j["VALOR MANUAL"], errors="coerce")
            calculado = pd.to_numeric(j["VALOR CALCULADO"], errors="coerce")
            j["DIFERENÇA"] = (calculado.fillna(0) - manual.fillna(0)).round(2)
            zero = (manual.fillna(0).abs() < 0.005) & (calculado.fillna(0).abs() < 0.005)
            j["RESULTADO"] = np.select(
                [(j["_merge"] == "left_only") & zero, (j["_merge"] == "right_only") & zero,
                 j["_merge"] == "left_only", j["_merge"] == "right_only", j["DIFERENÇA"].abs() >= 0.01],
                ["OK (R$ 0,00 só no manual)", "OK (R$ 0,00 só no cálculo)", "SÓ NO MANUAL", "SÓ NO CÁLCULO",
                 "VALOR DIFERENTE"], "OK")
            obs_calc = dict(zip(zip(calc["_k"], calc["_n"]), calc["TIPO PARCELA"].astype(str) + " " +
                                calc["SITUAÇÃO"].fillna("").astype(str)))
            j["OBSERVAÇÃO"] = [explica.get((k, n), "não encontrado no relatório do Pincel") if m == "left_only" else
                               (f"linha do Pincel que o manual não listou ({obs_calc.get((k, n), '')})"
                                if m == "right_only" else "")
                               for k, n, m in zip(j["_k"], j["_n"], j["_merge"])]
            j.insert(0, "ARQUIVO", arq.name)
            todas.append(j[["ARQUIVO", "RESULTADO"] + COMUNS + ["VALOR MANUAL", "VALOR CALCULADO", "DIFERENÇA",
                                                              "% PARCEIRO", "EXCEÇÃO PARCEIRO", "OBSERVAÇÃO"]])
            ok = j["RESULTADO"].str.startswith("OK")
            iguais, conta = int((j["RESULTADO"] == "OK").sum()), j["RESULTADO"].value_counts()
            t_man, t_calc = arred1(float(manual.sum())), arred1(float(total_id.get(pid, 0.0)))
            zeros = int(ok.sum()) - iguais
            res.append({"ARQUIVO": arq.name, "TIPO": "RELATÓRIO DO PARCEIRO",
                        "PARCEIRO": cad.at[pid, "NOME"] if pid else f"{nome.iloc[0] if len(nome) else '?'} "
                                                                   "(não está na PLANILHA GERAL)",
                        "PF / PJ": cad.at[pid, "PF / PJ"] if pid else None,
                        "SITUAÇÃO": cad.at[pid, "SITUAÇÃO"] if pid else "",
                        "LINHAS NO MANUAL": len(ref), "LINHAS NO CÁLCULO": len(calc), "LINHAS IGUAIS": iguais,
                        **{k: int(conta.get(k, 0)) for k in ("VALOR DIFERENTE", "SÓ NO MANUAL", "SÓ NO CÁLCULO")},
                        "TOTAL MANUAL": t_man, "TOTAL CALCULADO": t_calc, "DIFERENÇA": round(t_calc - t_man, 2),
                        "RESULTADO": "OK - VALORES BATEM" if ok.all() and abs(t_man - t_calc) < 0.01
                        else "DIVERGENTE",
                        "EXPLICAÇÃO": (f"{zeros} linha(s) de R$ 0,00 aparecem só de um lado (ver aba Conferência "
                                       "Linhas Parceiros)") if zeros else ""})
        elif tipo == "PLANILHA GERAL":
            res_pg = _conferir_planilha_geral(arq, ref, r, valor_planilha)
            res.extend(res_pg)
    if not res:
        return None, None
    return pd.DataFrame(res), (pd.concat(todas, ignore_index=True) if todas else pd.DataFrame())


def _conferir_planilha_geral(arq, ref, r, valor_planilha):
    """Uma linha por parceiro da PLANILHA GERAL manual: Valor manual x Valor calculado, com a explicação."""
    dados, cad = r["dados"], r["cadastro"]
    ref = ref[ref["NOME PARCEIRO"].notna()].copy()
    por_nome = {}
    for i, nomes in zip(cad.index, cad["_nomes"]):
        for n in nomes[:1]:
            por_nome.setdefault(n, []).append(i)
    linhas_pincel = dados[dados["_parceiro_id"].notna()].groupby("_parceiro_id")
    qtd = linhas_pincel.size().to_dict()
    pagos = dados[dados["PARCEIRO PAGO"]].groupby("_parceiro_id").size().to_dict()
    motivos = dados[dados["_parceiro_id"].notna() & ~dados["PARCEIRO PAGO"]].groupby("_parceiro_id")[
        "CLASSIFICAÇÃO PARCEIRO"].agg(lambda s: s.mode().iloc[0]).to_dict()
    rpa_total = r["resumo_parceiros"].set_index("_id")["TOTAL REPASSE"].to_dict() if len(r["resumo_parceiros"]) else {}
    valores_calc = {}
    for i, v in valor_planilha.items():
        if v > 0:
            valores_calc.setdefault(round(v, 2), []).append(cad.at[i, "NOME"])
    usados, saida = {}, []
    for _, lin in ref.iterrows():
        nb = nome_base(lin["NOME PARCEIRO"])
        ids = por_nome.get(nb, [])
        k = usados.get(nb, 0)
        usados[nb] = k + 1
        pid = ids[k] if k < len(ids) else (ids[0] if ids else None)
        manual = float(pd.to_numeric(lin["VALOR"], errors="coerce") or 0.0)
        calc = float(valor_planilha.get(pid, 0.0)) if pid else 0.0
        dif = round(calc - manual, 2)
        centavos = 0.01 <= abs(dif) <= 0.05
        if abs(dif) < 0.01:
            exp = "OK"
            if pid and cad.at[pid, "SITUAÇÃO"] == RPA and rpa_total.get(pid, 0) > 0:
                exp = f"OK (RPA: R$ {rpa_total[pid]:,.2f} vai na aba RPA (RH))".replace(",", "X").replace(
                    ".", ",").replace("X", ".")
        elif pid is None:
            exp = "nome não encontrado no cadastro"
        elif centavos and manual > 0 and calc > 0:
            exp = ("arredondamento: o manual somou o VALOR DO PARCEIRO do Pincel já arredondado linha a linha; o "
                   "programa soma as fórmulas sem arredondar, igual ao TOTAL do relatório do parceiro")
        elif manual > 0 and calc == 0 and not qtd.get(pid):
            exp = ("sem vendas no relatório do Pincel deste mês - o valor manual pode vir de outra fonte "
                   "(ex.: saldo do Prominas Acadêmico)")
        elif manual > 0 and calc == 0:
            exp = f"não pago pelo programa: {motivos.get(pid, cad.at[pid, 'MOTIVO SITUAÇÃO'])}"
        elif manual == 0 and calc > 0:
            exp = ("manual zerado, mas o parceiro está cadastrado (sem DISTRATO/POLO/RPA) e teve vendas no mês - "
                   "confirmar; se não deve receber, marque a situação na PLANILHA GERAL ou em D_Parceiros_Ajustes")
        else:
            outros = [n for n in valores_calc.get(round(manual, 2), []) if n != cad.at[pid, "NOME"]]
            if outros:
                exp = f"valor manual igual ao calculado para {outros[0]} (provável troca/cópia de linha no manual)"
            elif manual > calc:
                exp = "manual maior que o calculado (ajuste manual ou saldo de outra fonte?)"
            else:
                exp = "calculado maior que o manual (vendas que o manual não considerou?)"
            if pid and not pd.isna(cad.at[pid, "INÍCIO DA PARCERIA"]):
                exp += f" | parceria considerada a partir de {cad.at[pid, 'INÍCIO DA PARCERIA']:%d/%m/%Y}"
        saida.append({"ARQUIVO": arq.name, "TIPO": "PLANILHA GERAL", "PARCEIRO": lin["NOME PARCEIRO"],
                      "PF / PJ": lin.get("PF / PJ"), "SITUAÇÃO": cad.at[pid, "SITUAÇÃO"] if pid else "",
                      "LINHAS NO MANUAL": None, "LINHAS NO CÁLCULO": pagos.get(pid, 0) if pid else 0,
                      "LINHAS IGUAIS": None, "VALOR DIFERENTE": None, "SÓ NO MANUAL": None, "SÓ NO CÁLCULO": None,
                      "TOTAL MANUAL": round(manual, 2), "TOTAL CALCULADO": round(calc, 2), "DIFERENÇA": dif,
                      "RESULTADO": "OK - VALORES BATEM" if abs(dif) < 0.01 else (
                          "OK - CENTAVOS (ARREDONDAMENTO)" if centavos and manual > 0 and calc > 0 else "DIVERGENTE"),
                      "EXPLICAÇÃO": exp})
    return saida


# ========================================== POLOS ==========================================
def conferir_polos(r, pasta):
    """Compara linha a linha com relatórios de polo calculados à mão. (None, None) se não houver."""
    dados = r["dados"]
    polo_por_nome = dict(zip(ap(dados["POLO"], norm), dados["_polo_key"]))
    todas, resumo = [], []
    for arq in arquivos_referencia(pasta):
        try:
            ref = _ler_referencia(arq)
        except Exception as erro:
            print(f"[Conferência] Não foi possível ler {arq.name}: {erro}")
            continue
        if _tipo_referencia(ref) != "POLO":
            continue
        ref = ref[ref["MATRÍCULA"].notna() & ref["NOME (ALUNO)"].notna()].copy()
        chaves = ap(ref["POLO"], norm).map(polo_por_nome).dropna()
        if chaves.empty:
            print(f"[Conferência] {arq.name}: polo não encontrado no relatório do mês.")
            continue
        chave = chaves.mode().iloc[0]
        calc = dados[dados["_polo_key"] == chave].copy()
        for d in (ref, calc):
            d["_k"] = chave_linha(d)
            d["_n"] = d.groupby("_k").cumcount()
        lado_ref = ref[["_k", "_n"] + COMUNS + ["VALOR DO POLO"]].rename(
            columns={"VALOR DO POLO": "VALOR DO POLO (MANUAL)"})
        lado_calc = calc[["_k", "_n"] + COMUNS + ["VALOR DO POLO CALCULADO", "% POLO", "CLASSIFICAÇÃO VENDEDOR",
                                                  "PAPEL NA VENDA", "VIGÊNCIA DA REGRA"]]
        j = lado_ref.merge(lado_calc, on=["_k", "_n"], how="outer", suffixes=("", " (calc)"), indicator=True)
        for c in COMUNS:
            j[c] = j[c].where(j[c].notna(), j[f"{c} (calc)"])
        manual = pd.to_numeric(j["VALOR DO POLO (MANUAL)"], errors="coerce")
        calculado = pd.to_numeric(j["VALOR DO POLO CALCULADO"], errors="coerce")
        j["DIFERENÇA"] = (calculado.fillna(0) - manual.fillna(0)).round(2)
        j["RESULTADO"] = np.select(
            [j["_merge"] == "left_only", j["_merge"] == "right_only", j["DIFERENÇA"].abs() >= 0.01],
            ["SÓ NO MANUAL", "SÓ NO CÁLCULO", "VALOR DIFERENTE"], "OK")
        j.insert(0, "ARQUIVO", arq.name)
        j.insert(1, "POLO", chave)
        todas.append(j[["ARQUIVO", "POLO", "RESULTADO"] + COMUNS + ["VALOR DO POLO (MANUAL)",
                                                                  "VALOR DO POLO CALCULADO", "DIFERENÇA", "% POLO",
                                                                  "CLASSIFICAÇÃO VENDEDOR", "PAPEL NA VENDA",
                                                                  "VIGÊNCIA DA REGRA"]])
        iguais, conta = int((j["RESULTADO"] == "OK").sum()), j["RESULTADO"].value_counts()
        t_man, t_calc = round(float(manual.sum()), 2), round(float(calculado.sum()), 2)
        resumo.append({"ARQUIVO": arq.name, "POLO": chave, "LINHAS NA REFERÊNCIA": len(ref),
                       "LINHAS NO CÁLCULO": len(calc), "LINHAS IGUAIS": iguais,
                       **{k: int(conta.get(k, 0)) for k in ("VALOR DIFERENTE", "SÓ NO MANUAL", "SÓ NO CÁLCULO")},
                       "TOTAL REFERÊNCIA": t_man, "TOTAL CALCULADO": t_calc,
                       "RESULTADO": "OK - VALORES BATEM" if iguais == len(j) and abs(t_man - t_calc) < 0.01
                       else "DIVERGENTE"})
    return (pd.concat(todas, ignore_index=True), pd.DataFrame(resumo)) if resumo else (None, None)


__all__ = ["conferir_parceiros", "conferir_polos", "PAGOS"]
