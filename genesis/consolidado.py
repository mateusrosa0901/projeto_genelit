"""Planilha consolidada 'TOTAL REPASSE - <MÊS>-<ANO>.xlsx': resumo, Parceiros Digitais, RPA, conferências,
polos, pendências, memória de cálculo e regras."""
import datetime as dt
from pathlib import Path
from urllib.parse import quote

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font

import config
from .excel import AMARELO, BORDA, FMT_MOEDA, FMT_PCT, FONTE, NEGRITO, cabecalho, formatos, tabela
from .parceiros import C_NAO_CAD, RPA
from .util import MESES, moeda

COLUNA_ARQUIVO_SEPARADO = "ARQUIVO" + config.SUFIXO_RELATORIO_SEPARADO   # relatório de MAT-R e parcela 1


def _link(ws, linha, coluna, caminho):
    c = ws.cell(linha, coluna)
    c.hyperlink = quote(Path(caminho).as_posix())
    c.font = Font(name="Arial", size=10, color="FF0563C1", underline="single")


def aba_resumo(wb, r):
    ws = wb.active
    ws.title, ws.sheet_view.showGridLines = "Resumo", False
    df, ref = r["dados"], r["referencia"]
    ws["A1"] = "REPASSE PARCEIROS DIGITAIS E POLOS"
    ws["A1"].font = Font(name="Arial", size=14, bold=True, color="FF202020")
    pagto = df["DATA PAGTO"].dropna()
    ws["A2"] = (f"Referência: {MESES[ref.month - 1]}/{ref.year}  |  Relatório Pincel: {r['arquivo_receita'].name} "
                f"({len(df):,} pagamentos de {pagto.min():%d/%m/%Y} a {pagto.max():%d/%m/%Y})  |  "
                f"Gerado em {dt.datetime.now():%d/%m/%Y %H:%M}").replace(",", ".")
    ws["A2"].font = FONTE
    parc, val = r["resumo_parceiros"], r["valores_planilha"]
    cad = r["cadastro"]
    tipos = cad["PF / PJ"].map(lambda t: str(t).strip().upper() if isinstance(t, str) else "")
    pagos_planilha = {i: v for i, v in val.items() if v > 0}
    pj = sum(v for i, v in pagos_planilha.items() if tipos.get(i) == "PJ")
    outros = sum(v for i, v in pagos_planilha.items() if tipos.get(i) != "PJ")
    n_pj = sum(1 for i in pagos_planilha if tipos.get(i) == "PJ")
    rpa = parc[parc["SITUAÇÃO"].eq(RPA)] if len(parc) else parc
    total_rpa = float(rpa["TOTAL REPASSE"].sum()) if len(rpa) else 0.0
    nr, nc = r["nao_recebem"], r["nao_cadastrados"]
    polos = r.get("resumo_polos", pd.DataFrame())
    ativos = polos["STATUS"].str.startswith("ATIVO") if len(polos) else pd.Series(dtype=bool)
    desat = polos["STATUS"].str.startswith("DESATIVADO") if len(polos) else pd.Series(dtype=bool)

    def soma(base, filtro):
        return float(base.loc[filtro, "TOTAL REPASSE"].sum()) if len(base) else 0.0
    total_parc = sum(pagos_planilha.values()) + (0.0 if config.RPA_NA_PLANILHA_GERAL else total_rpa)
    linhas = [
        ("Pagamentos de alunos no mês (relatório do Pincel)", len(df), float(df["VALOR PAGTO"].sum())),
        ("Repasse PARCEIROS DIGITAIS - PLANILHA GERAL", len(pagos_planilha), sum(pagos_planilha.values())),
        ("     Pessoa Jurídica - pagamento com Nota Fiscal", n_pj, pj),
        ("     Pessoa Física / sem tipo na planilha - recibo", len(pagos_planilha) - n_pj, outros),
        ("Repasse PARCEIROS DIGITAIS - RPA (RH)" + (" (já incluído acima)" if config.RPA_NA_PLANILHA_GERAL else ""),
         len(rpa), total_rpa),
        ("TOTAL PARCEIROS DIGITAIS", len(parc), total_parc),
        ("Informativo: parceiros cadastrados que NÃO recebem (POLO, DISTRATO, venda no próprio polo...)",
         nr["NOME PARCEIRO (PLANILHA GERAL)"].nunique() if len(nr) else 0,
         float(nr["VALOR PELA REGRA"].sum()) if len(nr) else 0.0),
        ("Informativo: vendedores com EQUIPE de parceiro que não estão na PLANILHA GERAL (verificar se são "
         "parceiros novos - aba Vendedores Não Cadastrados)",
         int((nc["CLASSIFICAÇÃO"] == C_NAO_CAD).sum()) if len(nc) else 0,
         float(nc.loc[nc["CLASSIFICAÇÃO"] == C_NAO_CAD, "VALOR PELA REGRA"].sum()) if len(nc) else 0.0)]
    if config.GERAR_POLOS:
        linhas += [
            ("Repasse POLOS ATIVOS (a pagar)", int((ativos & (polos["TOTAL REPASSE"] > 0)).sum()) if len(polos) else 0,
             soma(polos, ativos)),
            ("Repasse POLOS DESATIVADOS (aguardando autorização de Expansão e Negócios)",
             int((desat & (polos["TOTAL REPASSE"] > 0)).sum()) if len(polos) else 0, soma(polos, desat)),
            ("TOTAL A PAGAR (parceiros digitais + polos ativos)", None, total_parc + soma(polos, ativos)),
            ("Pendências de polos para análise (aba Pendências Polos)", len(r.get("pendencias", [])), None)]
    for i, titulo in enumerate(["ITEM", "QUANTIDADE", "VALOR"], 1):
        cabecalho(ws.cell(4, i, titulo))
    for k, linha in enumerate(linhas, 5):
        for i, v in enumerate(linha, 1):
            c = ws.cell(k, i, v)
            c.font, c.border = NEGRITO if linha[0].startswith(("TOTAL", "Repasse")) else FONTE, BORDA
        ws.cell(k, 2).number_format, ws.cell(k, 3).number_format = "#,##0", FMT_MOEDA
    for col, w in zip("ABC", (92, 14, 20)):
        ws.column_dimensions[col].width = w
    k = 4 + len(linhas) + 2
    conf = r.get("conferencia_parceiros")
    if conf is not None and len(conf):
        ws.cell(k, 1, "CONFERÊNCIA COM AS PLANILHAS CALCULADAS À MÃO (pasta referencias)").font = NEGRITO
        k += 1
        rel = conf[conf["TIPO"].eq("RELATÓRIO DO PARCEIRO")]
        for _, c in rel.iterrows():
            texto = (f"{c['ARQUIVO']}: {c['RESULTADO']} - {int(c['LINHAS IGUAIS'])} de {int(c['LINHAS NO MANUAL'])} "
                     f"linhas iguais | manual {moeda(c['TOTAL MANUAL'])} x calculado {moeda(c['TOTAL CALCULADO'])}")
            ws.cell(k, 1, texto).font = NEGRITO if c["RESULTADO"] != "DIVERGENTE" else \
                Font(name="Arial", size=10, bold=True, color="FFC00000")
            k += 1
        pg = conf[conf["TIPO"].eq("PLANILHA GERAL")]
        if len(pg):
            ok = int(pg["RESULTADO"].str.startswith("OK").sum())
            texto = (f"PLANILHA GERAL: {ok} de {len(pg)} parceiros iguais (ou só centavos de arredondamento) | manual "
                     f"{moeda(pg['TOTAL MANUAL'].sum())} x calculado {moeda(pg['TOTAL CALCULADO'].sum())} "
                     f"(diferenças explicadas na aba Conferência Parceiros)")
            ws.cell(k, 1, texto).font = NEGRITO
            k += 1
        k += 1
    conf_p = r.get("conferencia_polos_resumo")
    if conf_p is not None and len(conf_p):
        ws.cell(k, 1, "CONFERÊNCIA DE POLOS (pasta referencias)").font = NEGRITO
        for _, c in conf_p.iterrows():
            k += 1
            texto = (f"{c['ARQUIVO']}: {c['RESULTADO']} - {c['LINHAS IGUAIS']} de {c['LINHAS NA REFERÊNCIA']} linhas "
                     f"iguais | manual {moeda(c['TOTAL REFERÊNCIA'])} x calculado {moeda(c['TOTAL CALCULADO'])}")
            ws.cell(k, 1, texto).font = NEGRITO if c["RESULTADO"] != "DIVERGENTE" else \
                Font(name="Arial", size=10, bold=True, color="FFC00000")
        k += 2
    notas = [
        "COMO LER",
        "- PLANILHA GERAL - <MÊS>-<ANO>.xlsx: a planilha de controle de sempre, com o 'Valor' de cada Parceiro "
        "Digital calculado. RPA, POLO, DISTRATO e NÃO ENVIAR REPASSE ficam zerados nela, como no manual. "
        "Responsáveis de polo que venderam como parceiro em outros polos e não estavam nela entram no fim.",
        "- RELATÓRIOS PARCEIROS DIGITAIS: um arquivo por parceiro, no mesmo layout do relatório feito à mão "
        "(VALOR DO PARCEIRO com fórmula e linha TOTAL). Os de RPA ficam na subpasta 'RPA (RH)'.",
        f"- Relatórios '{config.SUFIXO_RELATORIO_SEPARADO.strip(' -')}' (parceiros e polos): as parcelas MAT-R e 1 "
        "ficam num arquivo à parte, com os valores pela regra (D_Comissao); o outro arquivo traz as demais.",
        "- Parceiros Digitais: total por parceiro (a 'tabela dinâmica' do manual) e comparação com o Pincel.",
        "- RPA (RH): planilha para o RH (manual, pág. 10). Preencha o saldo da carteira do Prominas Acadêmico.",
        "- Parceiros Não Pagos: quem está na PLANILHA GERAL, vendeu no mês e não recebe, com o motivo.",
        "- Vendedores Não Cadastrados: quem vendeu e não está na PLANILHA GERAL. Se for um parceiro novo, "
        "inclua-o na PLANILHA GERAL (pasta dados) e rode de novo.",
        "- Cruzamento de Nomes: nomes do Pincel ligados à PLANILHA GERAL por CPF/CNPJ ou por nome parecido.",
        "- Conferência Parceiros: planilhas feitas à mão (pasta referencias) x cálculo, com a explicação.",
        "- Memória de Cálculo: cada pagamento do Pincel com o percentual e o motivo aplicados.",
        "- Regras / Exceções: percentuais de D_Comissao.csv e exceções de D_Excecoes.csv usados no cálculo."]
    for i, texto in enumerate(notas):
        ws.cell(k + i, 1, texto).font = NEGRITO if i == 0 else FONTE
    k += len(notas) + 1
    for a, b in r.get("info_arquivos", []):
        ws.cell(k, 1, f"{a}: {b}").font = FONTE
        k += 1


def aba_rpa(wb, rpa):
    ws = wb.create_sheet("RPA (RH)")
    if rpa.empty:
        ws["A1"] = "Nenhum parceiro RPA com repasse no mês."
        return
    cols = list(rpa.columns) + ["VALOR BRUTO DO REPASSE", f"DESCONTO INSS PATRONAL ({config.INSS_PATRONAL_RPA:.0%})"]
    larg = {"NOME DO PARCEIRO": 44, "CPF": 18}
    for i, col in enumerate(cols, 1):
        cabecalho(ws.cell(1, i, col))
        ws.column_dimensions[chr(64 + i)].width = larg.get(col, 22)
    ws.row_dimensions[1].height = 32
    for r, linha in enumerate(rpa.itertuples(index=False, name=None), 2):
        nome, cpf, pincel, saldo = linha
        valores = [nome, cpf, pincel, None if pd.isna(saldo) else saldo, f"=C{r}+D{r}",
                   f"=ROUND(E{r}*{config.INSS_PATRONAL_RPA},2)"]
        for i, v in enumerate(valores, 1):
            c = ws.cell(r, i, v)
            c.font, c.border = FONTE, BORDA
            if i >= 3:
                c.number_format = FMT_MOEDA
            if i == 4:
                c.fill = AMARELO
    ult = len(rpa) + 1
    c = ws.cell(ult + 1, 1, "TOTAL")
    c.font, c.border = NEGRITO, BORDA
    for i in range(3, 7):
        L = chr(64 + i)
        c = ws.cell(ult + 1, i, f"=SUM({L}2:{L}{ult})")
        c.font, c.border, c.number_format = NEGRITO, BORDA, FMT_MOEDA
    ws.cell(ult + 3, 1, "Manual, pág. 10: somar ao repasse do Pincel o saldo da carteira do parceiro no Prominas "
                        "Acadêmico (Transferências > Carteiras > CPF > saldo mensalidade - coluna amarela) e enviar "
                        "ao RH (dp1@, dp2@ e recursoshumanos@grupoprominas.com.br).").font = FONTE
    ws.freeze_panes = "A2"


def gerar_consolidado(r, destino):
    wb = Workbook()
    aba_resumo(wb, r)
    parc = r["resumo_parceiros"]
    ws = wb.create_sheet("Parceiros Digitais")
    if len(parc):
        tab = parc.drop(columns=["_id", "TOTAL SEM ARREDONDAR"]).copy()
        links = {"ARQUIVO": r["arquivos_parceiros"]}
        if config.SEPARAR_PARCELAS:
            links[COLUNA_ARQUIVO_SEPARADO] = r["arquivos_parceiros_separados"]
        for col, arquivos in links.items():
            tab[col] = [Path(a).name if a else "" for a in parc["_id"].map(arquivos).fillna("")]
        tabela(ws, tab, formatos(tab, {"CPF/CNPJ": "@", "Nº PLANILHA": "0"}),
               {"NOME PARCEIRO (PLANILHA GERAL)": 40, "NOME NO PINCEL": 34, "EXCEÇÃO": 40, "ARQUIVO": 40,
                COLUNA_ARQUIVO_SEPARADO: 48, "OBSERVAÇÃO (PLANILHA)": 24, "NOVO NA PLANILHA": 18, "SITUAÇÃO": 14,
                "PF / PJ": 9},
               ["QTD PAGAMENTOS", "REPASSE MATRÍCULAS", "REPASSE MENSALIDADES", "TOTAL REPASSE",
                "VALOR DO PARCEIRO (PINCEL)", "DIFERENÇA (CALC - PINCEL)"],
               destacar=lambda t: t["SITUAÇÃO"].eq(RPA))
        for col, arquivos in links.items():
            col_arq = list(tab.columns).index(col) + 1
            for k, i in enumerate(parc["_id"], 2):
                caminho = arquivos.get(i)
                if caminho:
                    _link(ws, k, col_arq, caminho)
    aba_rpa(wb, r["rpa"])
    for nome, df, larg, totais in [
        ("Parceiros Não Pagos", r["nao_recebem"], {"NOME PARCEIRO (PLANILHA GERAL)": 40, "NOME NO PINCEL": 34,
                                                    "MOTIVO": 70}, ["QTD PAGAMENTOS", "VALOR PELA REGRA",
                                                                     "VALOR DO PARCEIRO (PINCEL)"]),
        ("Vendedores Não Cadastrados", r["nao_cadastrados"], {"NOME PARCEIRO": 40, "MOTIVO": 60, "CLASSIFICAÇÃO": 34,
                                                              "EQUIPE": 22, "PARECIDO COM (PLANILHA GERAL)": 40},
         ["QTD PAGAMENTOS", "VALOR PELA REGRA", "VALOR DO PARCEIRO (PINCEL)"]),
        ("Cruzamento de Nomes", r["cruzamentos"], {"NOME NO PINCEL": 40, "NOME NA PLANILHA GERAL": 40,
                                                    "COMO FOI LIGADO": 30}, ["QTD PAGAMENTOS",
                                                                              "VALOR DO PARCEIRO CALCULADO"])]:
        ws = wb.create_sheet(nome)
        if df is not None and len(df):
            tabela(ws, df, formatos(df, {"CPF PARCEIRO": "@", "CPF/CNPJ NO PINCEL": "@"}), larg, totais)
        else:
            ws["A1"] = "Nada a listar."
    conf, linhas = r.get("conferencia_parceiros"), r.get("conferencia_parceiros_linhas")
    if conf is not None and len(conf):
        tab = conf.copy()
        tabela(wb.create_sheet("Conferência Parceiros"), tab, formatos(tab),
               {"ARQUIVO": 34, "PARCEIRO": 40, "EXPLICAÇÃO": 90, "RESULTADO": 20, "TIPO": 22},
               ["TOTAL MANUAL", "TOTAL CALCULADO", "DIFERENÇA"],
               destacar=lambda t: t["RESULTADO"].ne("OK - VALORES BATEM"))
    if linhas is not None and len(linhas):
        tabela(wb.create_sheet("Conferência Linhas Parceiros"), linhas, formatos(linhas, {"MATRÍCULA": "@"}),
               {"ARQUIVO": 34, "NOME (ALUNO)": 30, "NOME PARCEIRO": 30, "OBSERVAÇÃO": 50, "EXCEÇÃO PARCEIRO": 40},
               destacar=lambda t: t["RESULTADO"].ne("OK"))
    if config.GERAR_POLOS:
        polos = r["resumo_polos"]
        ws = wb.create_sheet("Repasse Polos")
        if len(polos):
            cols_arq = [c for c in ("ARQUIVO", COLUNA_ARQUIVO_SEPARADO) if c in polos.columns]
            tab = polos.drop(columns=["_polo_key"]).assign(
                **{c: [Path(a).name if a else "" for a in polos[c]] for c in cols_arq})
            tabela(ws, tab, formatos(tab), {"POLO": 34, "RESPONSÁVEL": 32, "E-MAIL": 30, "STATUS": 22, "ARQUIVO": 40,
                                            COLUNA_ARQUIVO_SEPARADO: 48},
                   [c for c in tab.columns if c in ("Graduação", "Segunda Graduação", "Tecnólogo", "Curso Técnico",
                                                    "Disciplina Isolada", "Outras", "TOTAL REPASSE",
                                                    "VALOR DO POLO (PINCEL)", "DIFERENÇA (CALC - PINCEL)",
                                                    "VALOR PAGO PELOS ALUNOS", "QTD PAGAMENTOS")],
                   destacar=lambda t: t["STATUS"].str.startswith("DESATIVADO"))
            for col in cols_arq:
                col_arq = list(tab.columns).index(col) + 1
                for k, caminho in enumerate(polos[col], 2):
                    if caminho:
                        _link(ws, k, col_arq, caminho)
        ws, pend = wb.create_sheet("Pendências Polos"), r["pendencias"]
        if len(pend):
            tabela(ws, pend, formatos(pend), {"TIPO": 36, "DESCRIÇÃO": 50, "O QUE FAZER": 80, "SUGESTÃO": 60})
        conf_l = r.get("conferencia_polos_linhas")
        if conf_l is not None and len(conf_l):
            tabela(wb.create_sheet("Conferência Polos"), conf_l, formatos(conf_l, {"MATRÍCULA": "@"}),
                   destacar=lambda t: t["RESULTADO"].ne("OK"))
    if config.GERAR_MEMORIA_DE_CALCULO:
        mem = r["memoria"]
        tabela(wb.create_sheet("Memória de Cálculo"), mem,
               formatos(mem, {"MATRÍCULA": "@", "PARCELA": "General", "LINHA ORIGEM": "0", "CPF PARCEIRO": "@"}),
               {"NOME (ALUNO)": 30, "CURSO": 30, "MOTIVO PARCEIRO": 40, "EXCEÇÃO PARCEIRO": 30, "POLO": 26,
                "POLO (CADASTRO)": 26, "NOME PARCEIRO": 30, "PARCEIRO DIGITAL": 30, "RESPONSÁVEL DO POLO": 28,
                "OBS PARCEIRO": 25, "VIGÊNCIA DA REGRA": 23, "CLASSIFICAÇÃO PARCEIRO": 30})
    ws, regras = wb.create_sheet("Regras"), r["regras"]
    tab = pd.DataFrame({"INÍCIO (MATRÍCULA)": regras["_ini"].dt.strftime("%d/%m/%Y"),
                        "FIM (MATRÍCULA)": regras["_fim"].dt.strftime("%d/%m/%Y"), "MODALIDADE": regras["modalidade"],
                        "TIPO PARCELA": regras["tipo_pagamento"], "PAPEL": regras["papel_beneficiario"],
                        "PERCENTUAL": regras["_pct"]})
    ult = tabela(ws, tab, {"PERCENTUAL": FMT_PCT}, {"INÍCIO (MATRÍCULA)": 14, "FIM (MATRÍCULA)": 14,
                                                    "MODALIDADE": 22, "TIPO PARCELA": 15, "PAPEL": 24,
                                                    "PERCENTUAL": 14})
    legenda = [
        "PAPEL: SELF_SALE = venda do responsável do polo ou sem parceiro (o polo recebe o percentual total);",
        "POLO_OUTRO_VENDEDOR = parte do polo quando outro parceiro/vendedor fez a matrícula;",
        "PARCEIRO_COMERCIAL = parte do Parceiro Digital.",
        "Parcelas numeradas (1, 2, 3...) = MENSALIDADE. MAT = matrícula (100% para o responsável pela matrícula).",
        f"Parceiro Digital: parcela(s) {', '.join(map(str, config.PARCELAS_SEM_REPASSE_PARCEIRO))} sem repasse; "
        f"'1ª mensalidade' das exceções = PARCELA {config.PARCELA_PRIMEIRA_MENSALIDADE}."]
    for i, texto in enumerate(legenda):
        ws.cell(ult + 2 + i, 1, texto).font = FONTE
    exc = r["excecoes"]
    if len(exc):
        t = exc[["parceiro", "cpf_cnpj", "papel", "polos", "modalidades", "modalidades_excluidas", "aplica_em",
                 "matricula_a_partir_de", "percentual", "observacao"]].copy()
        t.columns = [c.upper().replace("_", " ") for c in t.columns]
        t["PERCENTUAL"] = exc["_pct"].values
        tabela(wb.create_sheet("Exceções"), t, {"PERCENTUAL": FMT_PCT, "CPF CNPJ": "@"},
               {"PARCEIRO": 38, "OBSERVACAO": 90, "POLOS": 45, "CPF CNPJ": 17})
    wb.save(destino)
