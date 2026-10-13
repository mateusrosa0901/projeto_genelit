"""Leitura das entradas: relatório do Pincel, D_Polos, D_Comissao, D_Excecoes, PLANILHA GERAL e ajustes."""
import io
import re

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string

import config
from .util import (ap, datas, documento, exigir, ler_csv, nome_base, norm, pct, tokens, txt_int, uma_data,
                   valores, vazio)

COLUNAS_PINCEL = ["MATRÍCULA", "NOME (ALUNO)", "DATA DE MATRÍCULA", "NÍVEL", "CURSO", "POLO", "PARCELAS DO CURSO",
                  "VALOR DO CURSO", "PARCELA", "CÓDIGO PARCEIRO", "CPF PARCEIRO", "NOME PARCEIRO", "EQUIPE",
                  "VENCIMENTO", "DATA PAGTO", "VALOR PAGTO", "FORMA DE PAGAMENTO", "VALOR DO POLO",
                  "VALOR DO PARCEIRO", "SITUAÇÃO", "DATA SITUAÇÃO", "OBS PARCEIRO", "CH CURSO"]
COLUNAS_DATA = ["DATA DE MATRÍCULA", "VENCIMENTO", "DATA PAGTO", "DATA SITUAÇÃO"]
COLUNAS_VALOR = ["VALOR DO CURSO", "VALOR PAGTO", "VALOR DO POLO", "VALOR DO PARCEIRO"]


# =================================== PINCEL ===================================
def localizar_receita():
    """Relatório do Pincel: pasta 'entrada' (ou a pasta do projeto, como nas versões anteriores)."""
    for pasta in (config.PASTA_ENTRADA, config.PASTA_PROJETO):
        if (pasta / config.ARQ_RECEITA).exists():
            return pasta / config.ARQ_RECEITA
    cands = [p for pasta in (config.PASTA_ENTRADA, config.PASTA_PROJETO) if pasta.exists()
             for p in pasta.glob("relatorioReceitaPolo*") if p.suffix.lower() in (".xlsx", ".xls", ".csv")
             and not p.name.startswith("~$")]
    if not cands:
        raise FileNotFoundError(
            f"Relatório do Pincel não encontrado. Salve-o na pasta '{config.PASTA_ENTRADA}' como "
            f"'{config.ARQ_RECEITA}' (ou ajuste ARQ_RECEITA no config.py).")
    return max(cands, key=lambda p: p.stat().st_mtime)


def ler_receita():
    arq = localizar_receita()
    if arq.suffix.lower() == ".csv":
        df = ler_csv(arq)
    else:
        try:
            df = pd.read_excel(arq)
        except Exception as erro:  # o Pincel pode exportar HTML com extensão .xls (números no padrão brasileiro)
            try:
                bruto = arq.read_bytes()
                texto = None
                for cod in ("utf-8", "cp1252", "latin-1"):
                    try:
                        texto = bruto.decode(cod)
                        break
                    except UnicodeDecodeError:
                        continue
                df = pd.read_html(io.StringIO(texto), decimal=",", thousands=".")[0]
            except Exception:
                raise ValueError(f"Não foi possível ler {arq.name}: {erro}. "
                                 "Abra o arquivo no Excel e salve como .xlsx.") from erro
    por_nome = {norm(c): c for c in df.columns}  # aceita cabeçalho sem acento / com espaços extras
    df = df.rename(columns={por_nome[norm(c)]: c for c in COLUNAS_PINCEL if norm(c) in por_nome})
    exigir(df, COLUNAS_PINCEL, "no relatório do Pincel")
    df = df[COLUNAS_PINCEL].copy().astype(object)
    df.insert(0, "LINHA ORIGEM", np.arange(2, len(df) + 2))  # linha no Excel do Pincel
    df = df[~df["MATRÍCULA"].map(vazio)].reset_index(drop=True)
    for c in COLUNAS_DATA:
        df[c] = datas(df[c])
    for c in COLUNAS_VALOR:
        df[c] = valores(df[c])

    def parcela(v):
        t = txt_int(v)
        return int(t) if t.isdigit() else t.upper()
    df["MATRÍCULA"] = ap(df["MATRÍCULA"], lambda v: None if vazio(v) else txt_int(v))
    df["PARCELA"] = ap(df["PARCELA"], lambda v: None if vazio(v) else parcela(v))
    for c in ["CH CURSO", "PARCELAS DO CURSO", "CÓDIGO PARCEIRO"]:
        df[c] = pd.to_numeric(df[c], errors="coerce")
    for c in ["NOME (ALUNO)", "NÍVEL", "CURSO", "POLO", "NOME PARCEIRO", "EQUIPE", "FORMA DE PAGAMENTO", "SITUAÇÃO"]:
        df[c] = ap(df[c], lambda v: None if vazio(v) else str(v).strip())
    df["OBS PARCEIRO"] = ap(df["OBS PARCEIRO"], lambda v: None if vazio(v) else v)
    df["CPF PARCEIRO"] = ap(df["CPF PARCEIRO"], lambda v: None if vazio(v) else v)
    return arq, df


# ================================ TABELAS (CSV) ================================
def ler_polos():
    """Cadastro de polos; se o polo aparece mais de uma vez (responsável antigo + atual), vale o ATIVO."""
    arq = config.PASTA_DADOS / config.ARQ_POLOS
    p = ler_csv(arq)
    exigir(p, ["polo_key", "polo_pagamentos", "polo_nome", "estado", "cpf_responsavel", "nome_responsavel",
               "status_polo", "email"], f"em {arq.name}")
    p["_chave"] = ap(p["polo_key"], norm)
    p = p[p["_chave"].notna()].copy()
    p["_tokens"] = ap(p["polo_key"], tokens)
    p["_status"] = ap(p["status_polo"], norm)
    p["_status"] = p["_status"].where(p["_status"].notna(), "")
    p["_rank"] = p["_status"].map({"ATIVO": 0, "PARCEIRO_SEM_POLO": 1, "DESATIVADO": 2}).fillna(3)
    p["_doc"] = ap(p["cpf_responsavel"], documento)
    p["_nomes"] = ap(p["nome_responsavel"], lambda v: [] if vazio(v) else
                     [n for n in map(nome_base, re.split(r"/|\s-\s", str(v))) if n])
    p["_ordem"] = np.arange(len(p))
    principal = p.sort_values(["_rank", "_ordem"]).drop_duplicates("_chave").set_index("_chave")
    por_tokens = {}
    for chave, tk in zip(principal.index, principal["_tokens"]):
        por_tokens.setdefault(tk, chave)
    return p, principal, por_tokens


def ler_comissao():
    arq = config.PASTA_DADOS / config.ARQ_COMISSAO
    r = ler_csv(arq)
    r.columns = [c.lower() for c in r.columns]
    exigir(r, ["periodo_inicio", "periodo_fim", "modalidade", "tipo_pagamento", "papel_beneficiario", "percentual"],
           f"em {arq.name}")
    r["_mod"], r["_tipo"], r["_papel"] = (ap(r[c], norm) for c in ("modalidade", "tipo_pagamento",
                                                                    "papel_beneficiario"))
    r["_ini"], r["_fim"], r["_pct"] = datas(r["periodo_inicio"]), datas(r["periodo_fim"]), ap(r["percentual"], pct)
    invalida = r[["_mod", "_tipo", "_papel", "_ini", "_fim", "_pct"]].isna().any(axis=1)
    if invalida.any():
        print(f"[D_Comissao] ATENÇÃO: {int(invalida.sum())} linha(s) inválida(s) ignorada(s).")
    r = r[~invalida].reset_index(drop=True)
    r["_pct"] = r["_pct"].astype(float)
    return r


def ler_excecoes():
    arq = config.PASTA_DADOS / config.ARQ_EXCECOES
    if not arq.exists():
        print(f"[Exceções] {arq.name} não encontrado - nenhuma exceção aplicada.")
        return pd.DataFrame()
    e = ler_csv(arq)
    e.columns = [c.lower() for c in e.columns]
    exigir(e, ["parceiro", "cpf_cnpj", "papel", "polos", "modalidades", "modalidades_excluidas", "aplica_em",
               "matricula_a_partir_de", "percentual"], f"em {arq.name}")
    if "observacao" not in e.columns:
        e["observacao"] = None

    def lista(v):
        return [] if vazio(v) else [norm(x) for x in str(v).split("|") if norm(x)]
    e["_nome"], e["_doc"] = ap(e["parceiro"], nome_base), ap(e["cpf_cnpj"], documento)
    e["_papel"], e["_aplica"] = ap(e["papel"], norm), ap(e["aplica_em"], norm)
    for coluna, nova in (("polos", "_polos"), ("modalidades", "_mods"), ("modalidades_excluidas", "_mods_excl")):
        e[nova] = ap(e[coluna], lista)
    e["_desde"], e["_pct"] = datas(e["matricula_a_partir_de"]), ap(e["percentual"], pct)
    return e


# ============================ PLANILHA GERAL (cadastro) ============================
def _coluna(cabecalho, *nomes):
    alvo = {norm(n) for n in nomes}
    for i, v in enumerate(cabecalho, 1):
        if norm(v) in alvo:
            return i
    return None


def ler_planilha_geral(arq=None):
    """Lê a PLANILHA GERAL (controle de Parceiros Digitais) e devolve (cadastro, colunas).
    Cada linha com 'Nome Parceiro' preenchido é um parceiro; 'LINHA' é a linha no Excel."""
    arq = arq or config.PASTA_DADOS / config.ARQ_PLANILHA_GERAL
    if not arq.exists():
        raise FileNotFoundError(f"PLANILHA GERAL não encontrada: {arq}. Ela é o cadastro de Parceiros Digitais "
                                "(copie a planilha de controle do mês anterior para a pasta 'dados').")
    wb = load_workbook(arq, read_only=True, data_only=True)
    ws = wb.worksheets[0]
    linhas = list(ws.iter_rows(values_only=True))
    wb.close()
    if not linhas:
        raise ValueError(f"{arq.name} está vazia.")
    cab = list(linhas[0])
    col = {"Nº": _coluna(cab, "Nº", "N", "NO", "NUMERO"), "NOME": _coluna(cab, "Nome Parceiro", "NOME"),
           "CPF": _coluna(cab, "CPF"), "TIPO": _coluna(cab, "PF / PJ", "PF/PJ"), "EMAIL": _coluna(cab, "EMAIL"),
           "VALOR": _coluna(cab, "Valor"), "OBS": column_index_from_string(config.COLUNA_OBSERVACAO_PLANILHA)}
    falta = [k for k in ("NOME", "CPF", "TIPO", "EMAIL", "VALOR") if not col[k]]
    if falta:
        raise ValueError(f"{arq.name}: não encontrei as colunas {falta} na linha 1 (Nome Parceiro, CPF, PF / PJ, "
                         "EMAIL, Valor).")
    regs = []

    def valor(lin, k):
        i = col[k]
        return lin[i - 1] if i and i - 1 < len(lin) else None
    for n, lin in enumerate(linhas[1:], 2):
        def v(k, lin=lin):
            return valor(lin, k)
        if vazio(v("NOME")):
            continue
        regs.append({"ID": f"L{n}", "LINHA": n, "Nº": v("Nº"), "NOME": str(v("NOME")).strip(),
                     "CPF": None if vazio(v("CPF")) else str(v("CPF")).strip(),
                     "PF / PJ": None if vazio(v("TIPO")) else str(v("TIPO")).strip(),
                     "EMAIL": None if vazio(v("EMAIL")) else str(v("EMAIL")).strip(),
                     "VALOR NA PLANILHA": v("VALOR"),
                     "OBSERVAÇÃO": None if vazio(v("OBS")) else str(v("OBS")).strip()})
    return pd.DataFrame(regs).astype(object), col


def ler_ajustes_parceiros():
    """D_Parceiros_Ajustes.csv (opcional): apelidos no Pincel, documentos extras, data de início da parceria,
    situação forçada e parceiros que ainda não estão na PLANILHA GERAL."""
    arq = config.PASTA_DADOS / config.ARQ_AJUSTES_PARCEIROS
    colunas = ["nome_planilha", "apelidos_pincel", "cpf_cnpj_extra", "data_inicio_parceria", "situacao",
               "cpf_cnpj", "pf_pj", "email", "observacao"]
    if not arq.exists():
        return pd.DataFrame(columns=colunas)
    a = ler_csv(arq)
    a.columns = [c.lower() for c in a.columns]
    exigir(a, ["nome_planilha"], f"em {arq.name}")
    for c in colunas:
        if c not in a.columns:
            a[c] = None
    a = a[a["nome_planilha"].notna()].copy()
    a["_inicio"] = ap(a["data_inicio_parceria"], uma_data)
    return a
