"""Arquivos Excel: relatório de cada Parceiro Digital e PLANILHA GERAL (no layout dos arquivos feitos à mão),
relatório de cada polo e as tabelas do consolidado."""
import datetime as dt
from copy import copy

import numpy as np
import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import Cell
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as letra
from openpyxl.worksheet.dimensions import ColumnDimension
from openpyxl.worksheet.table import Table, TableStyleInfo

import config
from .util import texto_pct, vazio

# ------------------------------------ estilos ------------------------------------
FMT_MOEDA = r'_-"R$"\ * #,##0.00_-;\-"R$"\ * #,##0.00_-;_-"R$"\ * "-"??_-;_-@_-'
FMT_DATA, FMT_PCT = "mm-dd-yy", "0%"   # data curta do Excel (mostra dd/mm/aaaa no Brasil)
FONTE, NEGRITO = Font(name="Arial", size=10, color="FF202020"), Font(name="Arial", size=10, bold=True, color="FF202020")
FONTE_TOTAL = Font(name="Carlito", size=11, bold=True)
CINZA, BRANCO, VERDE, VERMELHO, AMARELO = (PatternFill("solid", fgColor=c) for c in
                                           ("FFD9D9D9", "FFFFFFFF", "FFB4E5A2", "FFFF0000", "FFFFF2CC"))
_FINO = Side(style="thin", color="FF000000")
BORDA = Border(left=_FINO, right=_FINO, top=_FINO, bottom=_FINO)
CENTRO, ESQ, MEIO = (Alignment(horizontal=h, vertical="center", wrap_text=True) for h in ("center", "left", None))


def celula(v):
    """Valor do pandas/numpy -> tipo aceito pelo Excel."""
    if isinstance(v, (list, tuple, set, dict)):
        return ", ".join(map(str, v)) if isinstance(v, (list, tuple, set)) else str(v)
    try:
        if v is None or pd.isna(v):
            return None
    except (TypeError, ValueError):
        pass
    if isinstance(v, pd.Timestamp):
        return v.to_pydatetime()
    return v.item() if isinstance(v, np.generic) else v


def carimbo(ws, font=FONTE, fill=None, align=None, fmt="General", border=BORDA):
    """Registra a combinação de estilo uma vez e devolve o 'carimbo' para aplicar rápido em milhares de células."""
    c = Cell(ws)
    c.font, c.number_format = font, fmt
    if border is not None:
        c.border = border
    if fill:
        c.fill = fill
    if align:
        c.alignment = align
    return c._style


def cabecalho(c):
    c.font, c.fill, c.border, c.alignment = NEGRITO, CINZA, BORDA, CENTRO


# ================= RELATÓRIO DO PARCEIRO DIGITAL (layout dos relatórios feitos à mão) =================
COLUNAS_RELATORIO_PARCEIRO = ["MATRÍCULA", "NOME (ALUNO)", "DATA DE MATRÍCULA", "NÍVEL", "CURSO", "POLO", "PARCELA",
                              "NOME PARCEIRO", "VENCIMENTO", "DATA PAGTO", "VALOR PAGTO", "FORMA DE PAGAMENTO",
                              "VALOR DO POLO", "VALOR DO PARCEIRO", "SITUAÇÃO", "DATA SITUAÇÃO", "OBS PARCEIRO",
                              "CH CURSO"]
_GS_MOEDA, _GS_DATA = "[$R$ -416]#,##0.00", "dd/mm/yyyy"
_GS_ALINHAMENTO = {"MATRÍCULA": "center", "NOME (ALUNO)": "left", "DATA DE MATRÍCULA": "center", "NÍVEL": "left",
                   "CURSO": "left", "POLO": "left", "PARCELA": "center", "NOME PARCEIRO": "left",
                   "VENCIMENTO": "center", "DATA PAGTO": "center", "VALOR PAGTO": "center",
                   "FORMA DE PAGAMENTO": "left", "VALOR DO POLO": "center", "VALOR DO PARCEIRO": "center",
                   "SITUAÇÃO": "right", "DATA SITUAÇÃO": "center", "OBS PARCEIRO": "left", "CH CURSO": "left"}
_GS_FORMATO = {"DATA DE MATRÍCULA": _GS_DATA, "VENCIMENTO": _GS_DATA, "DATA PAGTO": _GS_DATA,
               "VALOR PAGTO": _GS_MOEDA, "VALOR DO POLO": _GS_MOEDA, "VALOR DO PARCEIRO": _GS_MOEDA}
_GS_LARGURA = {"NOME (ALUNO)": 31.75, "DATA DE MATRÍCULA": 15.88, "PARCELA": 10.38, "FORMA DE PAGAMENTO": 18.88,
               "VALOR DO PARCEIRO": 20.25}


def formula_parceiro(linha_excel, p):
    """Fórmula do VALOR DO PARCEIRO como nos relatórios à mão: =K2*0.2, =K2*0.1, =K2 (100%) ou 0."""
    if p is None or np.isnan(p) or abs(p) < 1e-12:
        return 0
    if abs(p - 1) < 1e-12:
        return f"=K{linha_excel}"
    return f"=K{linha_excel}*{texto_pct(p)}"


def relatorio_parceiro(caminho, linhas, percentuais):
    """Um arquivo por Parceiro Digital, igual ao relatório calculado à mão (aba 'Página1', VALOR DO POLO oculto,
    VALOR DO PARCEIRO com fórmula e linha TOTAL). 'linhas' já vem ordenado pela DATA DE MATRÍCULA e
    'percentuais' traz o % do parceiro de cada linha (mesma ordem)."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Página1"
    ws.sheet_format.defaultRowHeight, ws.sheet_format.customHeight = 15.75, True
    ws.sheet_format.defaultColWidth = 12.63
    calibri = Font(name="Calibri", size=11, color="FF000000")
    cinza = PatternFill("solid", fgColor="FFC0C0C0")
    so_dir_baixo = Border(right=_FINO, bottom=_FINO)
    for i, col in enumerate(COLUNAS_RELATORIO_PARCEIRO, 1):
        c = ws.cell(1, i, col)
        c.font, c.fill = calibri, cinza
        c.border = BORDA if i == 1 else Border(right=_FINO, top=_FINO, bottom=_FINO)
        c.alignment = Alignment(horizontal="center", vertical="bottom")
        if col in _GS_LARGURA:
            ws.column_dimensions[letra(i)].width = _GS_LARGURA[col]
    estilos, estilo_data_sit = [], None
    for i, col in enumerate(COLUNAS_RELATORIO_PARCEIRO):
        borda = Border(left=_FINO, right=_FINO, bottom=_FINO) if i == 0 else so_dir_baixo
        alin = Alignment(horizontal=_GS_ALINHAMENTO[col], vertical="bottom")
        estilos.append(carimbo(ws, calibri, None, alin, _GS_FORMATO.get(col, "General"), borda))
        if col == "DATA SITUAÇÃO":
            estilo_data_sit = carimbo(ws, calibri, None, alin, _GS_DATA, borda)
    # VALOR DO PARCEIRO: verde quando o parceiro recebe, vermelho quando a linha fica zerada (como no manual)
    alin_n = Alignment(horizontal="center", vertical="bottom")
    verde_n, vermelho_n = (carimbo(ws, calibri, PatternFill("solid", fgColor=cor), alin_n, _GS_MOEDA, so_dir_baixo)
                           for cor in ("FF34A853", "FFEA4335"))
    i_parc, i_sit = COLUNAS_RELATORIO_PARCEIRO.index("VALOR DO PARCEIRO"), COLUNAS_RELATORIO_PARCEIRO.index(
        "DATA SITUAÇÃO")
    r = 1
    tabela_rel = linhas[COLUNAS_RELATORIO_PARCEIRO]
    for r, (linha, p) in enumerate(zip(tabela_rel.itertuples(index=False, name=None), percentuais), 2):
        for i, col in enumerate(COLUNAS_RELATORIO_PARCEIRO):
            v = celula(linha[i])
            if col == "MATRÍCULA" and v is not None:
                v = int(v) if str(v).isdigit() else v
            elif col == "VALOR DO POLO":
                v = None
            elif col == "VALOR DO PARCEIRO":
                v = formula_parceiro(r, p)
            c = ws.cell(r, i + 1, v)
            if i == i_parc:
                c._style = copy(vermelho_n if v == 0 else verde_n)
            else:
                c._style = copy(estilo_data_sit if (i == i_sit and v is not None) else estilos[i])
    ult = r
    for col_i, v in ((12, "TOTAL"), (13, None), (14, f"=SUM(N2:N{max(ult, 2)})")):
        c = ws.cell(ult + 1, col_i, v)
        c.font, c.border = Font(name="Arial", size=10, bold=True), BORDA
        c.alignment = Alignment(horizontal="center")
        if col_i == 14:
            c.number_format = _GS_MOEDA
    ws.column_dimensions["M"].hidden = True
    ws.auto_filter.ref = f"A1:R{max(ult, 2)}"
    wb.calculation.fullCalcOnLoad = True
    wb.save(caminho)


# ============================ PLANILHA GERAL (modelo do usuário) ============================
def planilha_geral(modelo, destino, cad, valores_por_id, referencia, colunas):
    """Abre a PLANILHA GERAL usada como cadastro e grava uma cópia com o 'Valor' do mês (mesma formatação).
    Parceiros que ainda não estão nela (D_Parceiros_Ajustes.csv) entram no fim, com o estilo da última linha."""
    wb = load_workbook(modelo)
    ws = wb.worksheets[0]
    cv = colunas["VALOR"]
    for i, linha in zip(cad.index, cad["LINHA"]):
        if not vazio(linha):
            ws.cell(int(linha), cv).value = float(valores_por_id.get(i, 0.0))
    cab = [c.value for c in ws[1]]
    ultima = max(int(x) for x in cad["LINHA"].dropna()) if cad["LINHA"].notna().any() else 1
    max_col = max(ws.max_column, colunas["OBS"])
    novos = cad[cad["NOVO"].astype(bool)]
    if len(novos):
        nums = [v for v in cad["Nº"] if isinstance(v, (int, float)) and not pd.isna(v)]
        prox = int(max(nums)) + 1 if nums else ultima
        modelo_linha = ultima
        por_titulo = {str(t).strip().upper(): i for i, t in enumerate(cab, 1) if t is not None}
        for k, (i, p) in enumerate(novos.iterrows(), 1):
            r = ultima + k
            for c in range(1, max_col + 1):
                origem = ws.cell(modelo_linha, c)
                destino_c = ws.cell(r, c)
                destino_c._style = copy(origem._style)
                v = origem.value
                # repete as fórmulas da linha (Data, CÓDIGO RECIBO, REFERENCIA DO MÊS: não dependem da linha)
                destino_c.value = v if isinstance(v, str) and v.startswith("=") else None
            ws.cell(r, colunas["Nº"] or 1).value = prox + k - 1
            ws.cell(r, colunas["NOME"]).value = p["NOME"]
            ws.cell(r, colunas["CPF"]).value = p["CPF"]
            ws.cell(r, colunas["TIPO"]).value = p["PF / PJ"]
            ws.cell(r, colunas["EMAIL"]).value = p["EMAIL"]
            ws.cell(r, cv).value = float(valores_por_id.get(i, 0.0))
            for titulo, valor in config.PLANILHA_GERAL_FIXOS.items():
                c = por_titulo.get(titulo.upper())
                if c:
                    ws.cell(r, c).value = valor
            ws.cell(r, colunas["OBS"]).value = f"INCLUÍDO PELO PROGRAMA ({p['ORIGEM']})"
            # colunas auxiliares do código do recibo (O/P = 0, 1, 2...): continua a sequência
            for c in range(1, max_col + 1):
                acima = ws.cell(r - 1, c).value
                if cab[c - 1] in (0, "0") and isinstance(acima, (int, float)):
                    ws.cell(r, c).value = acima + 1
        ultima += len(novos)
    # "Hoje" (base da REFERENCIA DO MÊS): data fixa no mês de referência, para não mudar ao reabrir o arquivo
    for i, t in enumerate(cab, 1):
        if str(t).strip().upper() == "HOJE":
            ws.cell(2, i).value = dt.datetime(referencia.year, referencia.month, 1)
    if ws.auto_filter.ref:
        ws.auto_filter.ref = f"A1:{letra(colunas['OBS'])}{ultima}"
    wb.calculation.fullCalcOnLoad = True
    wb.save(destino)


# ============================ RELATÓRIO DO POLO (layout de Açailândia) ============================
LAYOUT_POLO = {  # coluna: (largura, alinhamento, formato) - igual ao relatório manual de Açailândia
    "MATRÍCULA": (18, ESQ, "@"), "NOME (ALUNO)": (33.875, ESQ, "General"),
    "DATA DE MATRÍCULA": (19.125, CENTRO, FMT_DATA), "NÍVEL": (15.5, ESQ, "General"),
    "CURSO": (41.125, ESQ, "General"), "POLO": (14.5, ESQ, "General"), "PARCELA": (9.75, CENTRO, "0"),
    "NOME PARCEIRO": (29.25, ESQ, "General"), "VENCIMENTO": (15.5, ESQ, FMT_DATA), "DATA PAGTO": (15.5, ESQ, FMT_DATA),
    "VALOR PAGTO": (16.375, CENTRO, FMT_MOEDA), "FORMA DE PAGAMENTO": (24.5, ESQ, "General"),
    "VALOR DO POLO": (16.375, MEIO, FMT_MOEDA), "VALOR DO PARCEIRO": (18.75, CENTRO, FMT_MOEDA),
    "SITUAÇÃO": (14, ESQ, "General"), "DATA SITUAÇÃO": (17, ESQ, FMT_DATA), "OBS PARCEIRO": (15.25, ESQ, "General"),
    "CH CURSO": (12.25, ESQ, "0")}
COLUNAS_RELATORIO_POLO = list(LAYOUT_POLO)


def tabela_relatorio_polo(linhas):
    rel = pd.DataFrame({c: linhas[c] for c in COLUNAS_RELATORIO_POLO if c in linhas})
    rel["VALOR DO POLO"] = linhas["VALOR DO POLO CALCULADO"].values
    rel["VALOR DO PARCEIRO"] = 0.0   # o polo só vê o que ele recebe
    return rel[COLUNAS_RELATORIO_POLO]


def relatorio_polo(ws, rel):
    """Relatório no layout do arquivo manual 'POLO AÇAILÂNDIA - MA.xlsx' (anexo do e-mail do polo)."""
    cols = list(rel.columns)
    ws.sheet_view.showGridLines, ws.freeze_panes = False, "A2"
    for i, col in enumerate(cols, 1):
        c = ws.cell(1, i, col)
        cabecalho(c)
        c.number_format = "@" if col == "MATRÍCULA" else "General"
        d = ws.column_dimensions[letra(i)]  # a coluna inteira tem borda fina, como no modelo
        d.width, d.border = LAYOUT_POLO[col][0], BORDA
        d.alignment = Alignment(horizontal="center" if col == "PARCELA" else "left")
    resto = ColumnDimension(ws, index=letra(len(cols) + 1), width=9)
    resto.min, resto.max, resto.border, resto.alignment = len(cols) + 1, 16384, BORDA, Alignment(horizontal="left")
    ws.column_dimensions[letra(len(cols) + 1)] = resto
    ws.row_dimensions[1].height = 38.1
    estilos = [carimbo(ws, FONTE, BRANCO, LAYOUT_POLO[c][1], LAYOUT_POLO[c][2]) for c in cols]
    i_polo = cols.index("VALOR DO POLO")
    verde, vermelho = (carimbo(ws, FONTE, f, MEIO, FMT_MOEDA) for f in (VERDE, VERMELHO))
    for r, linha in enumerate(rel.itertuples(index=False, name=None), 2):
        for i, v in enumerate(linha):
            v = celula(v)
            if i == 0 and v is not None:
                v = str(v)  # MATRÍCULA como texto
            if i == i_polo:  # verde: o polo recebe; vermelho: o polo não recebe nesta linha
                ws.cell(r, i + 1, v)._style = copy(verde if v else vermelho)
            else:
                ws.cell(r, i + 1, v)._style = copy(estilos[i])
        ws.row_dimensions[r].height = 30
    ult = len(rel) + 1
    if len(rel):
        t = Table(displayName="DadosPolo", ref=f"A1:{letra(len(cols))}{ult}")
        t.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False,
                                          showRowStripes=True, showColumnStripes=False)
        ws.add_table(t)
    L = letra(i_polo + 1)   # linha em branco, rótulo "Repasse Total" e a soma
    for r, v in ((ult + 2, "Repasse Total"), (ult + 3, f"=SUM({L}2:{L}{max(ult, 2)})")):
        c = ws[f"{L}{r}"]
        c.value, c.font, c.border, c.alignment = v, FONTE_TOTAL, BORDA, Alignment(horizontal="center")
    ws[f"{L}{ult + 3}"].number_format = FMT_MOEDA
    for L in "KL":
        c = ws[f"{L}{ult + 3}"]
        if c.value is None:
            c.font, c.border, c.alignment = FONTE_TOTAL, BORDA, Alignment(horizontal="left")
    ws.row_dimensions[ult + 2].height = ws.row_dimensions[ult + 3].height = 15


# ===================================== TABELAS GENÉRICAS =====================================
def tabela(ws, df, fmts=None, larguras=None, totais=(), destacar=None, inicio=1):
    """Tabela com cabeçalho cinza, bordas, filtro, linhas destacadas em amarelo e linha de TOTAL."""
    fmts, larguras, cols = fmts or {}, larguras or {}, list(df.columns)
    for i, col in enumerate(cols, 1):
        cabecalho(ws.cell(inicio, i, col))
        w = larguras.get(col)
        if w is None:
            amostra = df[col].head(300).map(lambda v: len(str(v)) if celula(v) is not None else 0)
            w = min(max(len(col) * 0.9, (amostra.max() if len(amostra) else 8) * 1.1, 9), 60)
            w = max(w, 17) if fmts.get(col) == FMT_MOEDA else w
            w = max(w, 12) if col.startswith("QTD") else w
        ws.column_dimensions[letra(i)].width = w
    ws.row_dimensions[inicio].height = 32
    marcar = [bool(x) for x in destacar(df)] if destacar is not None else [False] * len(df)
    normal, amarelo = ([carimbo(ws, FONTE, f, None, fmts.get(c, "General")) for c in cols] for f in (None, AMARELO))
    for r, linha in enumerate(df.itertuples(index=False, name=None), inicio + 1):
        est = amarelo if marcar[r - inicio - 1] else normal
        for i, v in enumerate(linha):
            ws.cell(r, i + 1, celula(v))._style = copy(est[i])
    ult = inicio + len(df)
    ws.auto_filter.ref, ws.freeze_panes = f"A{inicio}:{letra(len(cols))}{max(ult, inicio)}", f"A{inicio + 1}"
    if totais and len(df):
        c = ws.cell(ult + 1, 1, "TOTAL")
        c.font, c.border = NEGRITO, BORDA
        for col in totais:
            i = cols.index(col) + 1
            c = ws.cell(ult + 1, i, f"=SUBTOTAL(9,{letra(i)}{inicio + 1}:{letra(i)}{ult})")
            c.font, c.border, c.number_format = NEGRITO, BORDA, fmts.get(col, FMT_MOEDA)
    return ult


def formatos(df, extras=None, datas=()):
    f = {c: FMT_MOEDA for c in df.columns if any(p in c for p in ("VALOR", "REPASSE", "DIFERENÇA", "DIF.", "SALDO",
                                                                    "INSS", "TOTAL", "MANUAL", "CALCULADO"))
         or c in ("Graduação", "Segunda Graduação", "Tecnólogo", "Curso Técnico", "Disciplina Isolada", "Outras")}
    for c in df.columns:
        if c.startswith("QTD") or c.startswith("LINHAS"):
            f[c] = "#,##0"
        if c.startswith("%"):
            f[c] = FMT_PCT
        if c in datas or c.startswith("DATA") or c in ("VENCIMENTO", "INÍCIO DA PARCERIA"):
            f[c] = FMT_DATA
    return {**f, **(extras or {})}
